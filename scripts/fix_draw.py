#!/usr/bin/env python3
"""모바일/Actions용: 이미 있는 회차 번호를 엑셀에서 수정 → draws/manifest 재생성.

Add(추가)와 달리 기존 회차만 허용합니다. add_draw.py / add-draw.yml 은 변경하지 않습니다.

사용 예:
  python scripts/fix_draw.py --draw-no 1236 --numbers 3,7,12,21,30,41 --bonus 15
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("pip install openpyxl")


def _parse_numbers(raw: str) -> list[int]:
    parts = [p.strip() for p in re.split(r"[,，\s]+", raw) if p.strip()]
    if len(parts) != 6:
        raise SystemExit(f"당첨 번호는 6개여야 합니다 (받은 개수: {len(parts)})")
    nums = [int(p) for p in parts]
    for n in nums:
        if n < 1 or n > 45:
            raise SystemExit(f"번호 범위 오류: {n} (1~45)")
    if len(set(nums)) != 6:
        raise SystemExit(f"당첨 번호에 중복이 있습니다: {nums}")
    return sorted(nums)


def _find_excel(repo: Path) -> Path:
    matches = list(repo.glob("*.xlsx"))
    matches = [p for p in matches if not p.name.startswith("~$")]
    if not matches:
        raise SystemExit("엑셀(.xlsx) 파일을 찾을 수 없습니다")
    for p in matches:
        if "당첨" in p.name or "lotto" in p.name.lower():
            return p
    return matches[0]


def _update_excel(path: Path, draw_no: int, win6: list[int], bonus: int) -> tuple[list[int], int]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    target_row = None
    old_win6: list[int] = []
    old_bonus = 0

    for row in ws.iter_rows(min_row=2, values_only=False):
        cell_b = row[1].value  # column B
        if cell_b is None:
            continue
        try:
            n = int(cell_b)
        except (TypeError, ValueError):
            continue
        if n != draw_no:
            continue
        target_row = row[1].row
        old_win6 = []
        for i in range(6):
            v = ws.cell(target_row, 3 + i).value
            old_win6.append(int(v) if v is not None else 0)
        bv = ws.cell(target_row, 9).value
        old_bonus = int(bv) if bv is not None else 0
        break

    if target_row is None:
        raise SystemExit(f"엑셀에 없는 회차입니다: {draw_no} (수정은 기존 회차만 가능)")

    for i, n in enumerate(win6):
        ws.cell(target_row, 3 + i, n)
    ws.cell(target_row, 9, bonus)
    wb.save(path)
    print(
        f"Excel update row={target_row} drawNo={draw_no} "
        f"from win6={old_win6} bonus={old_bonus} → win6={win6} bonus={bonus}"
    )
    return old_win6, old_bonus


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--draw-no", type=int, required=True)
    parser.add_argument("--numbers", required=True)
    parser.add_argument("--bonus", type=int, required=True)
    parser.add_argument("--repo-root", type=Path, default=Path("."))
    parser.add_argument(
        "--base-url",
        default="https://s73325475-cyber.github.io/lotto-data",
    )
    parser.add_argument("--version", type=int, default=3)
    parser.add_argument("--skip-convert", action="store_true")
    args = parser.parse_args()

    repo = args.repo_root.resolve()
    win6 = _parse_numbers(args.numbers)
    bonus = args.bonus
    if bonus < 1 or bonus > 45:
        raise SystemExit(f"보너스 범위 오류: {bonus}")
    if bonus in win6:
        raise SystemExit(f"보너스가 당첨 번호와 겹칩니다: {bonus}")

    excel = _find_excel(repo)
    print(f"Using excel: {excel.name}")
    _update_excel(excel, args.draw_no, win6, bonus)

    if args.skip_convert:
        return

    convert = repo / "convert_excel_to_draws.py"
    if not convert.is_file():
        raise SystemExit("convert_excel_to_draws.py 없음")

    dist = repo / "dist"
    dist.mkdir(exist_ok=True)
    cmd = [
        sys.executable,
        str(convert),
        str(excel),
        "--version",
        str(args.version),
        "--output-dir",
        str(dist),
        "--base-url",
        args.base_url,
    ]
    print("Running:", " ".join(cmd))
    subprocess.check_call(cmd)

    for name in (f"draws_v{args.version}.json", "manifest.json"):
        src = dist / name
        if src.is_file():
            (repo / name).write_bytes(src.read_bytes())
            print(f"Copied {name} → repo root")


if __name__ == "__main__":
    main()
